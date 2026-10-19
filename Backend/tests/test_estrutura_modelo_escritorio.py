"""A peca montada sobre o timbre segue a estrutura do modelo do escritorio.

Modelo padrao (timbre G. Trindade): I – PRELIMINARMENTE; II – MERITO com
subsecoes A), B)... e danos morais como subsecao do merito; III – LITIGANCIA
DE MA-FE; IV – DOS CALCULOS; V – AUTENTICIDADE; VI – PEDIDOS. Travessao '–'
(en dash) e sem secao separada de fundamentos.
"""

from __future__ import annotations

import base64
from io import BytesIO

from docx import Document

from App.services.contestacao_docx_builder import montar_docx_com_modelo

DADOS = {
    "numero_processo": "0001234-56.2026.5.06.0001",
    "autor": "Joao da Silva",
    "reu": "Empresa XYZ LTDA",
    "tipo_acao": "Trabalhista",
}


def _modelo_b64() -> str:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "TIMBRE DO ESCRITORIO"
    doc.add_paragraph("placeholder")
    out = BytesIO()
    doc.save(out)
    return base64.b64encode(out.getvalue()).decode("ascii")


def _minuta(**extra):
    base = {
        "tese_central": "Improcedencia.",
        "preliminares": "## I — PRELIMINARMENTE.\n\n### A) DA LIMITACAO DA CONDENACAO.\n\nTexto.",
        "merito": (
            "## II — MERITO.\n\n### II.A) DO CONTRATO.\n\nTexto do contrato.\n\n"
            "### II.B) DA JORNADA.\n\nTexto da jornada."
        ),
        "fundamentos": "Art. 818 CLT; Sumula 338 TST.",
        "pedidos": "Improcedencia total.",
    }
    base.update(extra)
    return base


def _titulos(minuta) -> list[str]:
    docx = montar_docx_com_modelo(_modelo_b64(), DADOS, minuta)
    assert docx is not None
    paragrafos = [p for p in Document(BytesIO(docx)).paragraphs if p.text.strip()]
    # Titulos sao os paragrafos inteiros em negrito + sublinhado.
    return [
        p.text.strip()
        for p in paragrafos
        if p.runs and all(r.bold and r.underline for r in p.runs if r.text.strip())
    ]


def test_estrutura_completa_segue_o_modelo_do_escritorio():
    titulos = _titulos(
        _minuta(
            # IA com o travessao errado: o builder normaliza pro do modelo.
            litigancia_ma_fe="## III — DA LITIGÂNCIA DE MÁ-FÉ.\n\nTexto da ma-fe.",
            danos_morais="## IV — DA IMPROCEDENCIA DO PEDIDO DE DANOS MORAIS.\n\nSem ato ilicito.",
            # Sem titulo: entra o padrao do modelo.
            calculos="Impugna-se os calculos da inicial.",
        )
    )
    romanos = [
        t
        for t in titulos
        if t.split(" ", 1)[0] in {"I", "II", "III", "IV", "V", "VI", "VII"}
    ]

    assert romanos[:6] == [
        "I – PRELIMINARMENTE.",
        "II – MERITO.",
        "III – DA LITIGÂNCIA DE MÁ-FÉ.",
        "IV – DOS CÁLCULOS.",
        "V – DA AUTENTICIDADE DOS DOCUMENTOS ORA ACOSTADOS.",
        "VI – DOS PEDIDOS.",
    ]


def test_danos_morais_vira_ultima_subsecao_do_merito_com_a_proxima_letra():
    titulos = _titulos(
        _minuta(
            danos_morais="## IV — DA IMPROCEDENCIA DO PEDIDO DE DANOS MORAIS.\n\nTexto."
        )
    )
    i_merito = titulos.index("II – MERITO.")
    i_proxima_secao = next(
        i
        for i, t in enumerate(titulos)
        if i > i_merito and t.startswith(("III", "IV", "V"))
    )
    subsecoes = titulos[i_merito + 1 : i_proxima_secao]

    # 'II.A)' do texto da IA vira 'A)' (letra simples, como no modelo)
    assert subsecoes == [
        "A) DO CONTRATO.",
        "B) DA JORNADA.",
        "C) DA IMPROCEDÊNCIA DO PEDIDO DE INDENIZAÇÃO POR DANOS MORAIS.",
    ]


def test_sem_secao_de_fundamentos_e_sem_travessao_longo():
    titulos = _titulos(_minuta())

    assert not any("FUNDAMENTOS" in t for t in titulos)
    assert not any("—" in t for t in titulos)


def test_secoes_opcionais_ausentes_renumeram_em_sequencia():
    titulos = _titulos(_minuta())
    romanos = [
        t for t in titulos if t.split(" ", 1)[0] in {"I", "II", "III", "IV", "V"}
    ]

    assert romanos[:4] == [
        "I – PRELIMINARMENTE.",
        "II – MERITO.",
        "III – DA AUTENTICIDADE DOS DOCUMENTOS ORA ACOSTADOS.",
        "IV – DOS PEDIDOS.",
    ]
