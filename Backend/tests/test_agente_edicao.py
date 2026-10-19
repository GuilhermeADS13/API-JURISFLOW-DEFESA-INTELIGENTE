"""Testes do agente de edicao no backend (piloto da migracao n8n -> backend)."""

from __future__ import annotations

import asyncio
import base64
import json
from io import BytesIO
from types import SimpleNamespace

import anthropic
import httpx2
import pytest
from docx import Document
from fastapi import HTTPException, Request

from App.routes import edicao as edicao_route
from App.services import agente_edicao

CAMPOS = {
    "nome": "Erica Cavalcante de Oliveira",
    "numero_processo": None,
    "valor_causa": None,
}
SUBSTITUICOES = [
    {
        "campo": "nome",
        "antigo": "Janaina Pereira da Silva Matos",
        "novo": "Erica Cavalcante de Oliveira",
        "ocorrencias_esperadas": 2,
    }
]


def _mensagem(dados: dict | None = None, stop_reason: str = "end_turn"):
    texto = json.dumps(
        dados
        if dados is not None
        else {"substituicoes": SUBSTITUICOES, "campos_ausentes": []}
    )
    return SimpleNamespace(
        content=[SimpleNamespace(type="text", text=texto)],
        stop_reason=stop_reason,
        usage=SimpleNamespace(input_tokens=1200, output_tokens=80),
    )


@pytest.fixture
def claude_fake(monkeypatch):
    """Troca o AsyncAnthropic por um fake que registra o pedido e devolve `resposta`."""
    estado = {"resposta": _mensagem(), "chamadas": []}

    class FakeMessages:
        async def create(self, **kwargs):
            estado["chamadas"].append(kwargs)
            if isinstance(estado["resposta"], Exception):
                raise estado["resposta"]
            return estado["resposta"]

    class FakeClient:
        def __init__(self, **kwargs):
            estado["client_kwargs"] = kwargs
            self.messages = FakeMessages()

    monkeypatch.setattr(agente_edicao.anthropic, "AsyncAnthropic", FakeClient)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-teste")
    monkeypatch.setenv("CLAUDE_MODEL", "claude-opus-5")
    return estado


def test_monta_pedido_com_schema_e_sem_temperature(claude_fake):
    resposta = asyncio.run(agente_edicao.gerar_substituicoes("texto", CAMPOS))

    assert resposta["status"] == "ok"
    assert resposta["substituicoes"] == SUBSTITUICOES
    pedido = claude_fake["chamadas"][0]
    assert pedido["model"] == "claude-opus-5"
    assert "temperature" not in pedido  # Opus 5 responde 400 se receber
    assert pedido["output_config"]["format"]["type"] == "json_schema"
    assert "Erica Cavalcante de Oliveira" in pedido["messages"][0]["content"]
    # SDK cuida do retry de 429/5xx/conexao.
    assert claude_fake["client_kwargs"]["max_retries"] == 2


def test_sem_chave_vira_erro_ia_sem_chamar_api(claude_fake, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY")
    resposta = asyncio.run(agente_edicao.gerar_substituicoes("texto", CAMPOS))

    assert resposta["status"] == "erro_ia"
    assert resposta["campos_ausentes"] == ["nome"]
    assert claude_fake["chamadas"] == []


@pytest.mark.parametrize(
    "erro, motivo",
    [
        (
            anthropic.APIStatusError(
                "overloaded",
                response=httpx2.Response(
                    529, request=httpx2.Request("POST", "https://x")
                ),
                body=None,
            ),
            "http_529",
        ),
        (
            anthropic.APIConnectionError(request=httpx2.Request("POST", "https://x")),
            "erro_conexao",
        ),
    ],
    ids=["http_529", "conexao"],
)
def test_erro_da_api_vira_erro_ia(claude_fake, erro, motivo):
    claude_fake["resposta"] = erro
    resposta = asyncio.run(agente_edicao.gerar_substituicoes("texto", CAMPOS))

    assert resposta["status"] == "erro_ia"
    assert resposta["_debug"]["fallback_motivo"] == motivo


@pytest.mark.parametrize("stop_reason", ["max_tokens", "refusal"])
def test_resposta_incompleta_vira_erro_ia(claude_fake, stop_reason):
    claude_fake["resposta"] = _mensagem(stop_reason=stop_reason)
    resposta = asyncio.run(agente_edicao.gerar_substituicoes("texto", CAMPOS))

    assert resposta["status"] == "erro_ia"
    assert resposta["substituicoes"] == []


# ── Rota com o engine padrao (agente no backend) ─────────────────────────────


def _payload():
    doc = Document()
    doc.add_paragraph("Reu: Janaina Pereira da Silva Matos, residente em Recife/PE.")
    doc.add_paragraph("A parte re Janaina Pereira da Silva Matos contesta os pedidos.")
    out = BytesIO()
    doc.save(out)
    return edicao_route.EdicaoContestacao.model_validate(
        {
            "arquivo_base_conteudo_base64": base64.b64encode(out.getvalue()).decode(
                "ascii"
            ),
            "arquivo_base_nome": "Contestacao.docx",
            "nome_novo": "Erica Cavalcante de Oliveira",
        }
    )


def _fake_request() -> Request:
    return Request(
        scope={
            "type": "http",
            "method": "POST",
            "path": "/api/editar-contestacao",
            "headers": [],
            "query_string": b"",
            "client": ("127.0.0.1", 0),
        }
    )


def _editar():
    return asyncio.run(
        edicao_route.editar_contestacao(
            request=_fake_request(),
            payload=_payload(),
            usuario={"id": "USR-TESTE", "nome": "Ana", "email": "a@a.com"},
        )
    )


def test_rota_usa_agente_do_backend_por_padrao(claude_fake, monkeypatch):
    monkeypatch.delenv("EDICAO_ENGINE", raising=False)

    async def n8n_nao_deve_ser_chamado(_dados):
        raise AssertionError("n8n chamado com o engine padrao")

    monkeypatch.setattr(
        edicao_route, "enviar_para_n8n_edicao", n8n_nao_deve_ser_chamado
    )

    resposta = _editar()

    assert resposta["ocorrencias_aplicadas"] == {"Janaina Pereira da Silva Matos": 2}
    doc = Document(BytesIO(base64.b64decode(resposta["arquivo_editado_base64"])))
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "Erica Cavalcante de Oliveira" in texto
    assert "Janaina" not in texto


def test_rota_devolve_503_quando_agente_falha(claude_fake, monkeypatch):
    monkeypatch.delenv("EDICAO_ENGINE", raising=False)
    claude_fake["resposta"] = _mensagem(stop_reason="max_tokens")

    with pytest.raises(HTTPException) as exc:
        _editar()
    assert exc.value.status_code == 503
